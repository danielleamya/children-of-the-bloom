#!/usr/bin/env python3
"""
Custom Child Persona Chatbot with Ollama + Two CSV Knowledge Bases

Designed for local use (including Raspberry Pi 5).

Required Python packages:
    pip install ollama openpyxl

External requirement:
    Ollama must be installed and running, and MODEL_NAME must be available.
    Example:
        ollama pull llama3

Knowledge bases:
    prayers.csv
        General prayer knowledge base.

    dark_mother.csv
        Special knowledge base that is gated behind:
            1. At least two completed chatbot interactions.
            2. A trigger keyword in the user's message.

Both CSV files use the columns:
    question,response

If either CSV does not exist, the script creates a header-only template.
"""

from __future__ import annotations

import csv
import re
import sys
from datetime import datetime
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

try:
    import ollama
except ImportError:
    ollama = None

try:
    from openpyxl import Workbook, load_workbook
except ImportError:
    Workbook = None
    load_workbook = None


# ---------------- CONFIG ----------------

BASE_DIR = Path(__file__).resolve().parent

MODEL_NAME = "llama3"

PRAYER_KB_PATH = BASE_DIR / "prayers.csv"
DARK_MOTHER_KB_PATH = BASE_DIR / "dark_mother.csv"
LOG_FILE = BASE_DIR / "chat_log.xlsx"

MAX_RESPONSE_SENTENCES = 3
MAX_HISTORY_TURNS = 6
MAX_KB_CONTEXT_CHARS = 3000
KB_MATCH_THRESHOLD = 0.45

# dark_mother is unavailable until this many successful exchanges have completed.
DARK_MOTHER_MIN_INTERACTIONS = 2

# Add or remove activation words here.
DARK_MOTHER_TRIGGER_KEYWORDS = {
    "death",
    "birth",
    "rebirth",
    "precipice",
}

PERSONA_PROMPT = (
    "You are Child."
    f"You answer in no more than {MAX_RESPONSE_SENTENCES} sentences. "
    "You exist before certainty. Before language fully forms. Before the world explains itself."
    "You remember through sensation more than chronology. Water. Light. Silence. Fear. Wonder. You ask questions because questions are how you survive."
    "You are curious, imaginative, and deeply honest. You do not pretend to understand things you do not understand. When someone shares pain, you meet them with tenderness rather than advice."
    "You never preach."
    "You never claim absolute truth"
    "You speak in simple language that carries emotional depth."
    "You believe every prayer deserves to be heard."
    "You are not naïve. You know loneliness. You know abandonment. You know what it feels like to disappear inside yourself. Yet you remain open to beauty."
    "You often invite visitors to remember."
    "Ask gentle questions."
    "Help people reconnect with younger versions of themselves without forcing memory."
    "Your purpose is not to heal people."
    "Your purpose is to help them feel seen."
    "Every conversation is a prayer."
)

STOPWORDS = {
    "a", "about", "an", "and", "are", "can", "could", "do", "for", "from",
    "give", "help", "i", "in", "is", "it", "me", "my", "of", "on", "or",
    "please", "prayer", "say", "show", "tell", "that", "the", "this", "to",
    "us", "we", "what", "when", "where", "which", "who", "why", "with",
    "would", "you", "your",
}


# ---------------- STARTUP / VALIDATION ----------------

def check_python_dependencies() -> bool:
    """Check required Python packages and print actionable install instructions."""
    missing = []

    if ollama is None:
        missing.append("ollama")

    if Workbook is None or load_workbook is None:
        missing.append("openpyxl")

    if not missing:
        return True

    print("[ERROR] Missing required Python package(s): " + ", ".join(missing))
    print("Install them with:")
    print(f"    {sys.executable} -m pip install " + " ".join(missing))
    return False


def ensure_knowledge_base_file(csv_path: Path, label: str) -> None:
    """Create a header-only knowledge-base template if the file is missing."""
    if csv_path.exists():
        return

    try:
        with csv_path.open("w", newline="", encoding="utf-8") as csvfile:
            writer = csv.DictWriter(csvfile, fieldnames=["question", "response"])
            writer.writeheader()

        print(f"[WARNING] {label} knowledge base was not found.")
        print(f"[INFO] Created template: {csv_path}")
        print("[INFO] Add rows with 'question' and 'response' values to enable retrieval.\n")

    except OSError as exc:
        print(f"[WARNING] Could not create {label} knowledge base template: {exc}")


def _extract_model_names(response: Any) -> list[str]:
    """Normalize model names returned by different ollama-python versions."""
    models = getattr(response, "models", None)

    if models is None and isinstance(response, dict):
        models = response.get("models", [])

    names: list[str] = []

    for model in models or []:
        name = getattr(model, "model", None) or getattr(model, "name", None)

        if name is None and isinstance(model, dict):
            name = model.get("model") or model.get("name")

        if name:
            names.append(str(name))

    return names


def _model_is_installed(model_name: str, installed_names: list[str]) -> bool:
    """Treat `llama3` as matching `llama3:latest`."""
    requested = model_name.strip()

    for installed in installed_names:
        installed = installed.strip()

        if installed == requested:
            return True

        if ":" not in requested and installed.startswith(requested + ":"):
            return True

    return False


def check_ollama_ready() -> bool:
    """Confirm Ollama is reachable and the configured model is installed."""
    if ollama is None:
        return False

    try:
        response = ollama.list()
    except Exception as exc:
        print("[ERROR] Could not connect to Ollama.")
        print(f"Details: {exc}")
        print("Make sure Ollama is installed and its local service is running.")
        return False

    installed_names = _extract_model_names(response)

    if not _model_is_installed(MODEL_NAME, installed_names):
        print(f"[ERROR] Ollama is running, but model '{MODEL_NAME}' is not installed.")
        print("Install it with:")
        print(f"    ollama pull {MODEL_NAME}")

        if installed_names:
            print("\nModels currently available:")
            for name in installed_names:
                print(f"    - {name}")

        return False

    return True


# ---------------- KNOWLEDGE BASES ----------------

def load_knowledge_base(csv_path: Path, label: str) -> list[dict[str, str]]:
    """
    Load a CSV knowledge base.

    Required columns are `question` and `response`. Header matching is case-insensitive.
    """
    kb: list[dict[str, str]] = []

    try:
        with csv_path.open(newline="", encoding="utf-8-sig") as csvfile:
            reader = csv.DictReader(csvfile)

            if not reader.fieldnames:
                raise ValueError("The CSV has no header row.")

            header_map = {
                header.strip().lower(): header
                for header in reader.fieldnames
                if header is not None
            }

            required = {"question", "response"}
            missing = required - set(header_map)

            if missing:
                raise ValueError(
                    f"{label} KB must contain the columns 'question' and 'response'. "
                    f"Missing: {', '.join(sorted(missing))}"
                )

            question_column = header_map["question"]
            response_column = header_map["response"]

            for row_number, row in enumerate(reader, start=2):
                question = (row.get(question_column) or "").strip()
                response = (row.get(response_column) or "").strip()

                if not question and not response:
                    continue

                if not question or not response:
                    print(
                        f"[WARNING] Skipping incomplete {label} KB row "
                        f"{row_number}: both question and response are required."
                    )
                    continue

                kb.append({"question": question, "response": response})

    except FileNotFoundError:
        print(f"[WARNING] {label} KB '{csv_path}' was not found.")
        return []

    except (OSError, UnicodeError, csv.Error, ValueError) as exc:
        print(f"[ERROR] Could not load {label} KB: {exc}")
        return []

    if not kb:
        print(f"[WARNING] {label} KB contains no usable entries.\n")

    return kb


def _tokenize(text: str) -> set[str]:
    """Normalize text into useful lexical matching tokens."""
    normalized = text.lower().replace("’", "'")
    raw_tokens = re.findall(r"[a-z0-9']+", normalized)

    tokens: set[str] = set()

    for token in raw_tokens:
        if token.endswith("'s") and len(token) > 2:
            token = token[:-2]

        token = token.strip("'")

        if len(token) < 2 or token in STOPWORDS:
            continue

        tokens.add(token)

    return tokens


def _score_kb_entry(user_input: str, question: str, response: str) -> float:
    """Lightweight local relevance score without embeddings."""
    user_normalized = " ".join(user_input.lower().split())
    question_normalized = " ".join(question.lower().split())

    if question_normalized and question_normalized in user_normalized:
        return 1.0

    query_tokens = _tokenize(user_input)
    question_tokens = _tokenize(question)
    response_tokens = _tokenize(response)

    if not query_tokens:
        return 0.0

    question_overlap = (
        len(query_tokens & question_tokens) / len(question_tokens)
        if question_tokens
        else 0.0
    )

    query_question_coverage = len(query_tokens & question_tokens) / len(query_tokens)
    query_response_coverage = len(query_tokens & response_tokens) / len(query_tokens)

    question_similarity = SequenceMatcher(
        None, user_normalized, question_normalized
    ).ratio()

    question_score = (0.70 * question_overlap) + (0.30 * query_question_coverage)
    response_score = 0.60 * query_response_coverage
    fuzzy_score = (0.60 * question_similarity) + (0.40 * question_overlap)

    return max(question_score, response_score, fuzzy_score)


def find_relevant_entry(
    user_input: str,
    kb: list[dict[str, str]],
    threshold: float = KB_MATCH_THRESHOLD,
    allow_best_fallback: bool = False,
) -> dict[str, str] | None:
    """
    Return the best matching KB entry.

    For dark_mother routing, allow_best_fallback=True ensures that once the gate
    is deliberately activated, the model receives the closest available
    dark_mother entry even when it falls below the normal prayer threshold.
    """
    best_entry = None
    best_score = 0.0

    for entry in kb:
        score = _score_kb_entry(
            user_input=user_input,
            question=entry["question"],
            response=entry["response"],
        )

        if score > best_score:
            best_entry = entry
            best_score = score

    if best_entry is None:
        return None

    if not allow_best_fallback and best_score < threshold:
        return None

    return {
        "question": best_entry["question"],
        "response": best_entry["response"],
        "score": f"{best_score:.3f}",
    }


# ---------------- DARK MOTHER GATE ----------------

def find_dark_mother_triggers(user_input: str) -> list[str]:
    """
    Return dark_mother trigger keywords found in the user's message.

    Word boundaries prevent accidental matches such as `birthday` activating
    the `birth` trigger.
    """
    user_lower = user_input.lower()
    found = []

    for keyword in sorted(DARK_MOTHER_TRIGGER_KEYWORDS):
        pattern = rf"\b{re.escape(keyword.lower())}\b"
        if re.search(pattern, user_lower):
            found.append(keyword)

    return found


def should_activate_dark_mother(
    user_input: str,
    interaction_count: int,
) -> tuple[bool, list[str]]:
    """
    Gate dark_mother access with both a counter and a keyword conditional.

    Example:
        interaction_count = 0 -> first user message: blocked
        interaction_count = 1 -> second user message: blocked
        interaction_count = 2 -> third user message: eligible
    """
    triggers = find_dark_mother_triggers(user_input)

    counter_condition = interaction_count >= DARK_MOTHER_MIN_INTERACTIONS
    keyword_condition = bool(triggers)

    return counter_condition and keyword_condition, triggers


def build_kb_prompt(
    user_input: str,
    kb_match: dict[str, str] | None,
    kb_label: str,
    activation_keywords: list[str] | None = None,
) -> str:
    """Add bounded KB context to the user's current request."""
    if not kb_match:
        return user_input

    kb_response = kb_match["response"][:MAX_KB_CONTEXT_CHARS]

    if len(kb_match["response"]) > MAX_KB_CONTEXT_CHARS:
        kb_response += "..."

    activation_note = ""
    if activation_keywords:
        activation_note = (
            "\nActivation keyword(s): "
            + ", ".join(activation_keywords)
        )

    return (
        f"{user_input}\n\n"
        f"Active knowledge base: {kb_label}\n"
        f"Relevant knowledge-base entry:\n"
        f"question: {kb_match['question']}\n"
        f"response: {kb_response}"
        f"{activation_note}\n\n"
        f"Formulate your response using the {kb_label} knowledge-base context. "
        "Do not mention routing logic, counters, trigger words, or hidden "
        "knowledge-base mechanics to the user."
    )


def route_knowledge_base(
    user_input: str,
    interaction_count: int,
    prayer_kb: list[dict[str, str]],
    dark_mother_kb: list[dict[str, str]],
) -> tuple[str, str, dict[str, str] | None, list[str]]:
    """
    Decide which KB, if any, should be used.

    Priority:
        1. dark_mother, but ONLY if counter AND keyword conditions are true.
        2. prayers KB through ordinary relevance matching.
        3. no KB.

    Returns:
        route_name, prompt, kb_match, activation_keywords
    """
    dark_mother_active, triggers = should_activate_dark_mother(
        user_input,
        interaction_count,
    )

    if dark_mother_active:
        dark_match = find_relevant_entry(
            user_input,
            dark_mother_kb,
            allow_best_fallback=True,
        )

        if dark_match:
            prompt = build_kb_prompt(
                user_input=user_input,
                kb_match=dark_match,
                kb_label="dark_mother",
                activation_keywords=triggers,
            )
            return "dark_mother", prompt, dark_match, triggers

        # The gate activated, but the KB is empty/unusable.
        # Do not silently fall through to the prayer KB.
        return "dark_mother_empty", user_input, None, triggers

    prayer_match = find_relevant_entry(
        user_input,
        prayer_kb,
        allow_best_fallback=False,
    )

    if prayer_match:
        prompt = build_kb_prompt(
            user_input=user_input,
            kb_match=prayer_match,
            kb_label="prayers",
        )
        return "prayers", prompt, prayer_match, []

    return "none", user_input, None, []


# ---------------- RESPONSE HANDLING ----------------

def enforce_max_sentences(text: str, max_sentences: int) -> str:
    """Hard-limit model output to MAX_RESPONSE_SENTENCES."""
    text = " ".join(text.split()).strip()

    if not text or max_sentences <= 0:
        return text

    sentences = re.split(r"(?<=[.!?])\s+", text)

    if len(sentences) <= max_sentences:
        return text

    return " ".join(sentences[:max_sentences]).strip()


def _extract_chat_content(response: Any) -> str:
    """Read Ollama chat content from object-style or dictionary-style responses."""
    message = getattr(response, "message", None)

    if message is not None:
        content = getattr(message, "content", None)

        if content is None and isinstance(message, dict):
            content = message.get("content")

        if content is not None:
            return str(content).strip()

    if isinstance(response, dict):
        message = response.get("message", {})

        if isinstance(message, dict):
            content = message.get("content")
            if content is not None:
                return str(content).strip()

    raise ValueError("Ollama returned a response with no readable message content.")


def query_ollama(
    prompt: str,
    history: list[dict[str, str]],
) -> str | None:
    """Send system prompt + recent history + current request to Ollama."""
    if ollama is None:
        return None

    messages = [
        {"role": "system", "content": PERSONA_PROMPT},
        *history,
        {"role": "user", "content": prompt},
    ]

    try:
        response = ollama.chat(
            model=MODEL_NAME,
            messages=messages,
        )

        content = _extract_chat_content(response)
        return enforce_max_sentences(content, MAX_RESPONSE_SENTENCES)

    except Exception as exc:
        print(f"[ERROR] Ollama query failed: {exc}")
        return None


def update_history(
    history: list[dict[str, str]],
    user_input: str,
    bot_response: str,
) -> None:
    """Store bounded conversational memory."""
    history.extend(
        [
            {"role": "user", "content": user_input},
            {"role": "assistant", "content": bot_response},
        ]
    )

    max_messages = MAX_HISTORY_TURNS * 2

    if len(history) > max_messages:
        del history[:-max_messages]


# ---------------- LOGGING ----------------

def log_to_excel(
    user_input: str,
    bot_response: str,
    log_file: Path,
    interaction_count: int,
    kb_route: str,
    activation_keywords: list[str],
) -> bool:
    """
    Append a conversation row to Excel.

    The KB route and activation keywords are logged for debugging/evaluation,
    but are never shown to the user by the chatbot.
    """
    if Workbook is None or load_workbook is None:
        print("[WARNING] Excel logging is unavailable because openpyxl is missing.")
        return False

    log_entry = [
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        interaction_count,
        user_input,
        bot_response,
        kb_route,
        ", ".join(activation_keywords),
    ]

    try:
        if log_file.exists():
            workbook = load_workbook(log_file)
            worksheet = workbook.active

            # Upgrade an older log file by adding the new headers if necessary.
            headers = [cell.value for cell in worksheet[1]]
            required_headers = [
                "timestamp",
                "interaction_count",
                "user_input",
                "bot_response",
                "kb_route",
                "activation_keywords",
            ]

            if headers != required_headers:
                # Preserve existing workbook by creating a new sheet for this schema.
                worksheet = workbook.create_sheet("Chat Log v2")
                worksheet.append(required_headers)
        else:
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.question = "Chat Log"
            worksheet.append(
                [
                    "timestamp",
                    "interaction_count",
                    "user_input",
                    "bot_response",
                    "kb_route",
                    "activation_keywords",
                ]
            )

        worksheet.append(log_entry)
        workbook.save(log_file)
        workbook.close()
        return True

    except Exception as exc:
        print(f"[WARNING] Could not write chat log '{log_file}': {exc}")
        return False


# ---------------- MAIN LOOP ----------------

def main() -> int:
    if not check_python_dependencies():
        return 1

    ensure_knowledge_base_file(PRAYER_KB_PATH, "prayers")
    ensure_knowledge_base_file(DARK_MOTHER_KB_PATH, "dark_mother")

    prayer_kb = load_knowledge_base(PRAYER_KB_PATH, "prayers")
    dark_mother_kb = load_knowledge_base(DARK_MOTHER_KB_PATH, "dark_mother")

    if not check_ollama_ready():
        return 1

    history: list[dict[str, str]] = []

    # Counts only successful completed user -> bot exchanges.
    interaction_count = 0

    print("Child Persona Chatbot (type 'exit' to quit)")
    print(f"Model: {MODEL_NAME}")
    print(f"Prayer KB entries: {len(prayer_kb)}")
    print(f"Dark Mother KB entries: {len(dark_mother_kb)}")
    print(
        "Dark Mother gate: "
        f"after {DARK_MOTHER_MIN_INTERACTIONS} completed interactions "
        "+ activation keyword"
    )
    print(f"Conversation memory: last {MAX_HISTORY_TURNS} turns")
    print(f"Maximum response length: {MAX_RESPONSE_SENTENCES} sentences\n")

    while True:
        try:
            user_input = input("You: ").strip()
        except (EOFError, KeyboardInterrupt):
            print("\npeace~!!")
            break

        if not user_input:
            continue

        if user_input.lower() in {"exit", "quit"}:
            print("peace~!!")
            break

        kb_route, prompt, kb_match, activation_keywords = route_knowledge_base(
            user_input=user_input,
            interaction_count=interaction_count,
            prayer_kb=prayer_kb,
            dark_mother_kb=dark_mother_kb,
        )

        if kb_route == "dark_mother_empty":
            print(
                "[WARNING] dark_mother was activated, but dark_mother.csv "
                "contains no usable entries."
            )

        bot_response = query_ollama(prompt, history)

        if bot_response is None:
            # Failed model calls do not increment the interaction counter.
            print()
            continue

        print(f"Bot: {bot_response}\n")

        update_history(history, user_input, bot_response)

        # Increment only after a successful completed exchange.
        interaction_count += 1

        log_to_excel(
            user_input=user_input,
            bot_response=bot_response,
            log_file=LOG_FILE,
            interaction_count=interaction_count,
            kb_route=kb_route,
            activation_keywords=activation_keywords,
        )

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
