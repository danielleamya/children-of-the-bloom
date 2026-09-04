"""
Rebuild audio cues from the Video Essay.docx line structure.

Display text follows the essay (paragraphs / sentences, names, casing).
Start/end times are aligned from Whisper word timestamps.

Usage (from Frontend/Ars):
  python scripts/rebuild_cues_from_docx.py
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

try:
    import docx
except ImportError:
    import subprocess

    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx", "-q"])
    import docx

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "assets" / "docs"
LOOPS_JSON = DOCS / "loops.json"
DOCX_PATH = DOCS / "Child_ren of the Bloom - Video Essay.docx"
XLSX = DOCS / "conversational_template.xlsx"
WHISPER_DIR = DOCS / "whisper_raw"
CHILD01_RAW = DOCS / "Child01_whisper_raw.json"

VIDEO_BY_KEY = {
    "Child01": "Video01Sound.mp4",
    "Child02": "Video02Sound.mp4",
    "Nana03": "Video03Sound.mp4",
    "Phytoplankton": "Video04Sound.mp4",
    "ChildMe05": "Video05Sound.mp4",
    "MassiveAlgal06": "Video06Sound.mp4",
    "EarthVideo07": "Video07Sound.mp4",
    "CosmicMother08": "Video08Sound.mp4",
}

ABBREV = re.compile(r"\b(mr|mrs|ms|dr|st|jr|sr|vs|etc)\.$", re.I)


def normalize(text: str) -> str:
    text = text.lower().replace("\u2019", "'").replace("\u2018", "'")
    text = re.sub(r"[^a-z0-9'\s]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def split_sentences(block: str) -> list[str]:
    text = re.sub(r"\s+", " ", (block or "").strip())
    if not text:
        return []
    parts: list[str] = []
    buf = ""
    for i, ch in enumerate(text):
        buf += ch
        if ch not in ".?!":
            continue
        nxt = text[i + 1] if i + 1 < len(text) else ""
        if nxt and not nxt.isspace():
            continue
        if ABBREV.search(buf.strip()):
            continue
        trimmed = buf.strip()
        if trimmed:
            parts.append(trimmed)
        buf = ""
        while i + 1 < len(text) and text[i + 1].isspace():
            i += 1
    tail = buf.strip()
    if tail:
        parts.append(tail)
    return parts


def parse_docx_sections(path: Path) -> tuple[list[dict], str | None]:
    document = docx.Document(str(path))
    paras: list[str] = []
    for p in document.paragraphs:
        t = re.sub(r"\s+", " ", (p.text or "").replace("\u00a0", " ").strip())
        if t:
            paras.append(t)

    sections: list[dict] = []
    closing_lines: list[str] = []
    in_closing = False
    current_paras: list[str] = []

    def is_question(text: str) -> bool:
        return (
            (text.startswith("What") or text.startswith("Who") or text.startswith("How will"))
            and "?" in text
        )

    def commit(paras_in: list[str], question: str) -> None:
        lines: list[str] = []
        for para in paras_in:
            lines.extend(split_sentences(para))
        prayer = "\n".join(lines).strip()
        if prayer and question:
            sections.append(
                {
                    "prayer": prayer,
                    "leadingQuestion": question,
                    "lines": lines,
                }
            )

    for p in paras:
        low = p.lower()
        if low.startswith("a child") and "prayer" in low:
            continue
        if low.startswith("video instructions"):
            in_closing = True
            continue
        if low.startswith("video"):
            continue
        if in_closing:
            closing_lines.append(p)
            continue
        if is_question(p):
            if current_paras:
                commit(current_paras, p)
                current_paras = []
            continue
        current_paras.append(p)

    closing = "\n".join(closing_lines).strip() or None
    return sections, closing


def load_whisper_words(key: str) -> list[dict]:
    if key == "Child01" and CHILD01_RAW.is_file():
        path = CHILD01_RAW
    else:
        path = WHISPER_DIR / f"{key}_whisper_raw.json"
    if not path.is_file():
        return []
    data = json.loads(path.read_text(encoding="utf-8"))
    words: list[dict] = []
    for seg in data.get("segments") or []:
        for w in seg.get("words") or []:
            token = (w.get("word") or "").strip()
            if not token:
                continue
            words.append(
                {
                    "word": token,
                    "start": float(w.get("start") or 0),
                    "end": float(w.get("end") or 0),
                }
            )
        if not seg.get("words") and (seg.get("text") or "").strip():
            words.append(
                {
                    "word": seg["text"].strip(),
                    "start": float(seg.get("start") or 0),
                    "end": float(seg.get("end") or 0),
                }
            )
    return words


def flatten_words(words: list[dict]) -> list[tuple[str, float, float]]:
    flat: list[tuple[str, float, float]] = []
    for w in words:
        for token in normalize(w["word"]).split():
            if token:
                flat.append((token, w["start"], w["end"]))
    return flat


def align_lines(lines: list[str], words: list[dict]) -> list[dict]:
    flat = flatten_words(words)
    cursor = 0
    cues: list[dict] = []
    n = len(flat)
    duration = flat[-1][2] if flat else 0.0

    for li, line in enumerate(lines):
        target = normalize(line).split()
        if not target:
            continue

        start_idx = None
        search_limit = min(n, cursor + 60)
        for i in range(cursor, max(cursor + 1, search_limit)):
            if i >= n:
                break
            if flat[i][0] == target[0] or (
                len(target[0]) >= 4 and flat[i][0].startswith(target[0][:4])
            ):
                ok = True
                for j, tok in enumerate(target[: min(4, len(target))]):
                    if i + j >= n:
                        ok = False
                        break
                    fw = flat[i + j][0]
                    if not (
                        fw == tok
                        or (len(tok) >= 3 and fw.startswith(tok[:3]))
                        or (len(fw) >= 3 and tok.startswith(fw[:3]))
                    ):
                        ok = False
                        break
                if ok:
                    start_idx = i
                    break

        if start_idx is None:
            # Place after previous cue; estimate duration from length.
            prev_end = cues[-1]["end"] if cues else 0.0
            est = max(1.2, min(8.0, 0.32 * max(1, len(target))))
            start_t = prev_end
            end_t = min(duration or start_t + est, start_t + est)
            cues.append(
                {
                    "text": line,
                    "start": round(start_t, 2),
                    "end": round(end_t, 2),
                    "matched": False,
                }
            )
            continue

        i = start_idx
        t = 0
        end_t = flat[start_idx][2]
        while t < len(target) and i < n:
            fw = flat[i][0]
            tok = target[t]
            if fw == tok or (len(tok) >= 3 and fw.startswith(tok[:3])) or (
                len(fw) >= 3 and tok.startswith(fw[:3])
            ):
                end_t = flat[i][2]
                t += 1
                i += 1
            elif i + 1 < n and (
                flat[i + 1][0] == tok
                or (len(tok) >= 3 and flat[i + 1][0].startswith(tok[:3]))
            ):
                i += 1
            else:
                t += 1

        cues.append(
            {
                "text": line,
                "start": round(flat[start_idx][1], 2),
                "end": round(end_t, 2),
                "matched": True,
            }
        )
        cursor = max(i, start_idx + 1)

    # Monotonic starts
    for i in range(1, len(cues)):
        if cues[i]["start"] < cues[i - 1]["start"]:
            cues[i]["start"] = cues[i - 1]["start"]
        if cues[i - 1]["end"] > cues[i]["start"]:
            cues[i - 1]["end"] = cues[i]["start"]

    for c in cues:
        if len(c["text"]) <= 28:
            c["nowrap"] = True
        c.pop("matched", None)
    return cues


def main() -> int:
    if not DOCX_PATH.is_file():
        print(f"Missing {DOCX_PATH}", file=sys.stderr)
        return 1

    sections, closing = parse_docx_sections(DOCX_PATH)
    print(f"Parsed {len(sections)} sections from essay")

    payload = json.loads(LOOPS_JSON.read_text(encoding="utf-8"))
    loops = payload.get("loops") or []

    for i, loop in enumerate(loops):
        key = loop.get("videoKey") or f"loop{i}"
        essay = sections[i] if i < len(sections) else None
        if not essay:
            print(f"[{i}] {key}: no essay section")
            continue

        lines = essay["lines"]
        words = load_whisper_words(key)
        cues = align_lines(lines, words) if words else [
            {"text": line, "start": float(i), "end": float(i) + 1.5}
            for i, line in enumerate(lines)
        ]
        matched_note = f"{len(cues)} lines"
        loop["prayer"] = "\n".join(lines)
        loop["leadingQuestion"] = essay["leadingQuestion"]
        loop["video"] = VIDEO_BY_KEY.get(key, loop.get("video"))
        loop["audioSync"] = True
        loop["audioCues"] = cues
        print(f"[{i}] {key}: {matched_note} (doc text + whisper timing)")
        # Show name check for Nana
        if key == "Nana03":
            hit = next((c["text"] for c in cues if "Casilda" in c["text"]), None)
            print(f"    Casilda line: {hit!r}")

    if closing and payload.get("closing"):
        payload["closing"]["prayer"] = closing

    LOOPS_JSON.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote {LOOPS_JSON}")

    try:
        import openpyxl
    except ImportError:
        return 0

    if XLSX.is_file():
        wb = openpyxl.load_workbook(XLSX)
        ws = wb.active
        loop_i = 0
        for row in range(2, ws.max_row + 1):
            if not ws.cell(row, 1).value:
                continue
            if not ws.cell(row, 2).value:
                if payload.get("closing"):
                    ws.cell(row, 1).value = payload["closing"].get("prayer")
                continue
            if loop_i >= len(loops):
                break
            loop = loops[loop_i]
            ws.cell(row, 1).value = loop.get("prayer")
            ws.cell(row, 2).value = loop.get("leadingQuestion")
            ws.cell(row, 3).value = VIDEO_BY_KEY.get(loop.get("videoKey"), loop.get("videoKey"))
            loop_i += 1
        wb.save(XLSX)
        print(f"Updated {XLSX}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
