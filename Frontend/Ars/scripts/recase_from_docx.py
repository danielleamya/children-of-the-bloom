"""
Rebuild loop prayers/cues from Whisper speech + Video Essay.docx casing.

- Spoken wording comes from Whisper (recording differs from the page).
- Capitalization / proper-noun style comes from the Word essay.

Usage (from Frontend/Ars):
  python scripts/recase_from_docx.py
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

try:
    import docx
except ImportError:
    import subprocess

    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx", "-q"])
    import docx

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "assets" / "docs"
LOOPS_JSON = DOCS / "loops.json"
DOCX_PATH = DOCS / "Child_ren of the Bloom - Video Essay.docx"
XLSX = DOCS / "conversational_template.xlsx"
WHISPER_DIR = DOCS / "whisper_raw"
CHILD01_RAW = DOCS / "Child01_whisper_raw.json"

VIDEO_BY_KEY = {
    "Child01": "Video01Sound.mp4",
    "Child02": "Video02Sound.mp4",
    "Nana03": "Video03Sound.mp4",
    "Phytoplankton": "Video04Sound.mp4",
    "ChildMe05": "Video05Sound.mp4",
    "MassiveAlgal06": "Video06Sound.mp4",
    "EarthVideo07": "Video07Sound.mp4",
    "CosmicMother08": "Video08Sound.mp4",
}

TOKEN_RE = re.compile(r"[A-Za-z0-9]+(?:['’][A-Za-z0-9]+)?|[^\sA-Za-z0-9]+", re.UNICODE)

# Whisper often mangles proper nouns / set phrases; prefer essay forms.
HARD_PHRASE_FIXES: list[tuple[re.Pattern[str], str]] = [
    (re.compile(r"\bkissel\s+de\s+ben+ers?\b", re.I), "Casilda Benners"),
    (re.compile(r"\bkisseldeben+ers?\b", re.I), "Casilda Benners"),
    (re.compile(r"\bcasilda\s+benners\b", re.I), "Casilda Benners"),
    (re.compile(r"\balgor\s*-?\s*booms?\b", re.I), "algal blooms"),
    (re.compile(r"\balgorbooms?\b", re.I), "algal blooms"),
    (re.compile(r"\bchildren of the blue\b", re.I), "Children of the Bloom"),
    (re.compile(r"\bchildren of the bloom\b", re.I), "Children of the Bloom"),
    (re.compile(r"\bblack liberated reality\b", re.I), "Black Liberated Reality"),
    (re.compile(r"\ba liberated reality\b", re.I), "a Liberated Reality"),
    (re.compile(r"\bmother oif the sea\b", re.I), "Mother of the Sea"),
    (re.compile(r"\bmother of the sea\b", re.I), "Mother of the Sea"),
    (re.compile(r"\bmother of the air\b", re.I), "Mother of the Air"),
    (re.compile(r"\bmother of the earth\b", re.I), "Mother of the Earth"),
    (re.compile(r"\bmother of the sun\b", re.I), "Mother of the Sun"),
    (re.compile(r"\bgreat depression\b", re.I), "Great Depression"),
    (re.compile(r"\bpuerto rico\b", re.I), "Puerto Rico"),
    (re.compile(r"\bnew york\b", re.I), "New York"),
    (re.compile(r"\bsan juan\b", re.I), "San Juan"),
    (re.compile(r"\bst\.?\s*croix\b", re.I), "St. Croix"),
    (re.compile(r"\bst\.?\s*thomas\b", re.I), "St. Thomas"),
    (re.compile(r"\bon board the philadelphia\b", re.I), "onboard the Philadelphia"),
    (re.compile(r"\bonboard the philadelphia\b", re.I), "onboard the Philadelphia"),
    (re.compile(r"\bportal'?s keeper\b", re.I), "portals keeper"),
    (re.compile(r"\buncle slick\b", re.I), "uncle Slick"),
    (re.compile(r"\blook\s+deaf\b", re.I), "look death"),
]

# If Whisper dropped a known opening, restore from essay.
PREFIX_FIXES_BY_KEY: dict[str, list[tuple[str, str]]] = {
    "Child02": [
        (
            "entire life at the end of the world",
            "I've lived my entire life at the end of the world",
        ),
    ],
}


def apply_hard_phrase_fixes(text: str) -> str:
    out = text
    for pattern, repl in HARD_PHRASE_FIXES:
        out = pattern.sub(repl, out)
    return out


def apply_prefix_fixes(text: str, key: str) -> str:
    stripped = text.lstrip()
    low = stripped.lower()
    for needle, full in PREFIX_FIXES_BY_KEY.get(key, []):
        if low.startswith(needle.lower()):
            # Preserve any leading whitespace from Whisper segment.
            lead = text[: len(text) - len(stripped)]
            # If already has the prefix words, leave it.
            if low.startswith(full.lower()):
                return text
            rest = stripped[len(needle) :]
            return f"{lead}{full}{rest}"
    return text


def extract_essay_phrases(essay: str) -> list[str]:
    """Multi-word proper phrases from the essay for soft replacement."""
    phrases: list[str] = []
    for m in re.finditer(
        r"\b(?:[A-Z][A-Za-z'’.-]*)(?:\s+(?:of|the|and|de|da|del|van|von))?(?:\s+[A-Z][A-Za-z'’.-]*)+\b",
        essay or "",
    ):
        phrase = m.group(0).strip()
        if len(phrase.split()) >= 2:
            phrases.append(phrase)
    # Prefer longer phrases first.
    phrases.sort(key=lambda p: (-len(p.split()), -len(p)))
    # Dedupe case-insensitively
    seen: set[str] = set()
    out: list[str] = []
    for p in phrases:
        k = p.lower()
        if k in seen:
            continue
        seen.add(k)
        out.append(p)
    return out


def soft_replace_essay_phrases(spoken: str, phrases: list[str]) -> str:
    """Replace near-matches of essay phrases using token overlap (names/places)."""
    tokens = spoken.split()
    if not tokens:
        return spoken
    norm_tokens = [normalize_token(t) for t in tokens]
    used = [False] * len(tokens)

    for phrase in phrases:
        p_toks = [normalize_token(t) for t in phrase.split() if normalize_token(t)]
        if len(p_toks) < 2:
            continue
        n = len(p_toks)
        best_i = -1
        best_score = 0.0
        for i in range(0, max(0, len(norm_tokens) - n + 1)):
            if any(used[i : i + n]):
                continue
            window = norm_tokens[i : i + n]
            if not any(window):
                continue
            hits = sum(
                1
                for a, b in zip(window, p_toks)
                if a and (a == b or a.startswith(b[:4]) or b.startswith(a[:4]))
            )
            score = hits / n
            # Exact-ish only — fuzzy last-token swaps (Sea→Sun) are too risky.
            if window[-1] != p_toks[-1]:
                continue
            if score >= 0.9 and score > best_score:
                best_score = score
                best_i = i
        if best_i >= 0:
            # Replace window with essay phrase tokens, keep trailing punctuation from last spoken token.
            last = tokens[best_i + n - 1]
            trail = re.sub(r"^.*?([.,!?;:]*)$", r"\1", last) if re.search(r"[.,!?;:]$", last) else ""
            replacement = phrase + trail
            tokens[best_i : best_i + n] = [replacement] + [""] * (n - 1)
            for j in range(best_i, best_i + n):
                used[j] = True
    return re.sub(r"\s+", " ", " ".join(t for t in tokens if t)).strip()


def normalize_token(tok: str) -> str:
    t = tok.lower().replace("’", "'")
    t = re.sub(r"[^a-z0-9']+", "", t)
    return t


def tokenize(text: str) -> list[str]:
    return [m.group(0) for m in TOKEN_RE.finditer(text or "")]


def is_word(tok: str) -> bool:
    return bool(re.search(r"[A-Za-z0-9]", tok))


def parse_docx_sections(path: Path) -> tuple[list[dict], str | None]:
    document = docx.Document(str(path))
    paras: list[str] = []
    for p in document.paragraphs:
        t = re.sub(r"\s+", " ", (p.text or "").replace("\u00a0", " ").strip())
        if t:
            paras.append(t)

    sections: list[dict] = []
    closing_lines: list[str] = []
    in_closing = False
    current_lines: list[str] = []

    def is_question(text: str) -> bool:
        return (
            (text.startswith("What") or text.startswith("Who") or text.startswith("How will"))
            and "?" in text
        )

    def commit(prayer_lines: list[str], question: str) -> None:
        prayer = "\n".join(prayer_lines).strip()
        if prayer and question:
            sections.append({"prayer": prayer, "leadingQuestion": question})

    for p in paras:
        low = p.lower()
        if low.startswith("a child") and "prayer" in low:
            continue
        if low.startswith("video instructions"):
            in_closing = True
            continue
        if low.startswith("video"):
            continue
        if in_closing:
            closing_lines.append(p)
            continue
        if is_question(p):
            if current_lines:
                commit(current_lines, p)
                current_lines = []
            continue
        current_lines.append(p)

    closing = "\n".join(closing_lines).strip() or None
    return sections, closing


def build_case_forms(original: str) -> tuple[dict[str, list[str]], set[str]]:
    """Map normalized token -> essay surface forms; track mid-sentence proper nouns."""
    forms: dict[str, list[str]] = {}
    proper: set[str] = set()
    sentence_start = True
    for tok in tokenize(original):
        if not is_word(tok):
            if tok in ".?!":
                sentence_start = True
            continue
        key = normalize_token(tok)
        if not key:
            continue
        forms.setdefault(key, []).append(tok)
        # Capitalized mid-sentence in the essay => treat as a proper noun.
        if (
            not sentence_start
            and tok[:1].isupper()
            and not tok.isupper()
            and key not in {"i", "i'm"}
        ):
            proper.add(key)
        if tok.isupper() and len(tok) > 1:
            proper.add(key)
        sentence_start = False
    return forms, proper


def fuzzy_norm(tok: str) -> str:
    n = normalize_token(tok)
    if n.endswith("'s") and len(n) > 3:
        return n[:-2]
    return n


def pick_case(
    norm: str,
    forms: dict[str, list[str]],
    proper: set[str],
    sentence_start: bool,
) -> str | None:
    # First-person pronoun is always capitalized in the essay.
    if norm in {"i", "i'm"}:
        cands = forms.get(norm) or forms.get("i'm" if norm == "i" else "i") or []
        for c in cands:
            if c[:1].isupper():
                return c
        return "I" if norm == "i" else "I'm"

    # Litany "maybe" is capitalized in the essay only at line starts.
    if norm == "maybe" and not sentence_start:
        return "maybe"

    candidates = forms.get(norm) or forms.get(fuzzy_norm(norm))
    if not candidates:
        if norm.endswith("s") and norm[:-1] in forms:
            candidates = forms[norm[:-1]]
        elif norm + "s" in forms:
            candidates = forms[norm + "s"]
    if not candidates:
        return None

    if not sentence_start:
        for c in candidates:
            if c[:1].islower():
                return c
        sample = candidates[0]
        if sample.isupper() and len(sample) > 1:
            return sample
        if norm in proper or fuzzy_norm(norm) in proper:
            return sample
        # Essay only capitalized this at sentence starts — keep it lower mid-line.
        return sample.lower()

    for c in candidates:
        if c[:1].isupper():
            return c
    return candidates[0]


def recase_text(
    spoken: str,
    forms: dict[str, list[str]],
    proper: set[str],
    *,
    initial_sentence_start: bool = True,
) -> str:
    tokens = tokenize(spoken)
    out: list[str] = []
    sentence_start = initial_sentence_start

    for tok in tokens:
        if not is_word(tok):
            out.append(tok)
            if tok in ".?!":
                sentence_start = True
            continue

        norm = normalize_token(tok)
        chosen = pick_case(norm, forms, proper, sentence_start)
        if chosen is None:
            if sentence_start:
                surface = tok[:1].upper() + tok[1:] if tok else tok
            else:
                surface = tok.lower() if not (tok.isupper() and len(tok) > 1) else tok
            if tok.isupper() and len(tok) > 1:
                alts = forms.get(norm) or []
                caps = [a for a in alts if a.isupper()]
                surface = caps[0] if caps else tok
            out.append(surface)
        else:
            surface = apply_case_pattern(tok, chosen)
            if sentence_start and surface and surface[:1].islower():
                surface = surface[:1].upper() + surface[1:]
            out.append(surface)
        sentence_start = False

    parts: list[str] = []
    for i, tok in enumerate(out):
        if i == 0:
            parts.append(tok)
            continue
        prev = parts[-1]
        if is_word(prev) and is_word(tok):
            parts.append(" ")
            parts.append(tok)
        elif is_word(tok) and prev in "\"'(":
            parts.append(tok)
        elif not is_word(tok) and tok in ".,!?;:%)]}":
            parts.append(tok)
        elif is_word(tok):
            parts.append(" ")
            parts.append(tok)
        else:
            parts.append(tok)
    text = "".join(parts)
    text = re.sub(r"\s+", " ", text).strip()
    text = re.sub(r"\s+([.,!?;:])", r"\1", text)
    return text


def apply_case_pattern(spoken: str, model: str) -> str:
    """Apply casing pattern from model onto spoken characters (best effort)."""
    if model.islower():
        return spoken.lower()
    if model.isupper() and len(model) > 1:
        return spoken.upper()
    if model[:1].isupper() and model[1:].islower():
        return spoken[:1].upper() + spoken[1:].lower()
    s = list(spoken)
    mi = 0
    for i, ch in enumerate(s):
        if not ch.isalpha():
            continue
        while mi < len(model) and not model[mi].isalpha():
            mi += 1
        if mi >= len(model):
            s[i] = ch.lower()
            continue
        s[i] = ch.upper() if model[mi].isupper() else ch.lower()
        mi += 1
    return "".join(s)


def load_whisper_segments(key: str) -> list[dict]:
    if key == "Child01" and CHILD01_RAW.is_file():
        path = CHILD01_RAW
    else:
        path = WHISPER_DIR / f"{key}_whisper_raw.json"
    if not path.is_file():
        return []
    data = json.loads(path.read_text(encoding="utf-8"))
    return data.get("segments") or []


def cues_from_recased_segments(
    segments: list[dict],
    forms: dict[str, list[str]],
    proper: set[str],
    *,
    key: str = "",
    essay_phrases: list[str] | None = None,
) -> tuple[str, list[dict]]:
    cues: list[dict] = []
    lines: list[str] = []
    prev_ended_sentence = True
    phrases = essay_phrases or []
    for seg in segments:
        raw = (seg.get("text") or "").strip()
        if not raw:
            continue
        raw = apply_prefix_fixes(raw, key)
        raw = apply_hard_phrase_fixes(raw)
        if phrases:
            raw = soft_replace_essay_phrases(raw, phrases)
        text = recase_text(
            raw,
            forms,
            proper,
            initial_sentence_start=prev_ended_sentence,
        )
        # Re-apply name fixes in case recase altered token boundaries.
        text = apply_hard_phrase_fixes(text)
        if text and prev_ended_sentence and text[:1].islower():
            text = text[:1].upper() + text[1:]
        if not text:
            continue
        lines.append(text)
        cue = {
            "text": text,
            "start": round(float(seg.get("start") or 0), 2),
            "end": round(float(seg.get("end") or 0), 2),
        }
        if (
            len(text) <= 28
            and text.endswith(".")
            or re.search(r"black liberated reality", text, re.I)
        ):
            cue["nowrap"] = True
        cues.append(cue)
        prev_ended_sentence = bool(re.search(r"[.!?]$", text))
    prayer = "\n".join(lines)
    return prayer, cues


def main() -> int:
    if not DOCX_PATH.is_file():
        print(f"Missing {DOCX_PATH}", file=sys.stderr)
        return 1
    if not LOOPS_JSON.is_file():
        print(f"Missing {LOOPS_JSON}", file=sys.stderr)
        return 1

    sections, closing = parse_docx_sections(DOCX_PATH)
    print(f"Parsed {len(sections)} essay sections from docx")
    if len(sections) < 8:
        print("WARNING: expected 8 sections", file=sys.stderr)

    payload = json.loads(LOOPS_JSON.read_text(encoding="utf-8"))
    loops = payload.get("loops") or []
    # Preserve editor wrap setting when present.
    saved_wrap = payload.get("wrapFraction")
    if saved_wrap is None:
        for loop in loops:
            if loop.get("wrapFraction") is not None:
                saved_wrap = loop.get("wrapFraction")
                break

    for i, loop in enumerate(loops):
        key = loop.get("videoKey") or f"loop{i}"
        essay = sections[i] if i < len(sections) else None
        original = (essay or {}).get("prayer") or loop.get("prayer") or ""
        forms, proper = build_case_forms(original)
        # Also include leading question tokens for terms like Black Liberated Reality
        if essay and essay.get("leadingQuestion"):
            q_forms, q_proper = build_case_forms(essay["leadingQuestion"])
            for k, v in q_forms.items():
                forms.setdefault(k, []).extend(v)
            proper |= q_proper

        essay_phrases = extract_essay_phrases(original)
        # Always keep critical name phrases available even if regex missed them.
        for forced in (
            "Casilda Benners",
            "Black Liberated Reality",
            "Children of the Bloom",
            "Mother of the Sun",
            "St. Croix",
            "St. Thomas",
            "Puerto Rico",
            "Great Depression",
        ):
            if forced.lower() not in {p.lower() for p in essay_phrases}:
                essay_phrases.insert(0, forced)

        segments = load_whisper_segments(key)
        if not segments and loop.get("audioCues"):
            # Fall back to recasing existing cues in place.
            print(f"[{i}] {key}: recasing existing cues from docx")
            new_cues = []
            lines = []
            prev_ended = True
            for cue in loop["audioCues"]:
                raw = apply_hard_phrase_fixes(cue.get("text") or "")
                raw = soft_replace_essay_phrases(raw, essay_phrases)
                text = recase_text(
                    raw,
                    forms,
                    proper,
                    initial_sentence_start=prev_ended,
                )
                text = apply_hard_phrase_fixes(text)
                nc = {**cue, "text": text}
                new_cues.append(nc)
                lines.append(text)
                prev_ended = bool(re.search(r"[.!?]$", text))
            loop["prayer"] = "\n".join(lines)
            loop["audioCues"] = new_cues
        elif segments:
            prayer, cues = cues_from_recased_segments(
                segments,
                forms,
                proper,
                key=key,
                essay_phrases=essay_phrases,
            )
            loop["prayer"] = prayer
            loop["audioCues"] = cues
            print(f"[{i}] {key}: {len(cues)} cues recased from Whisper + docx")
        else:
            print(f"[{i}] {key}: no whisper data; skipped")

        loop["video"] = VIDEO_BY_KEY.get(key, loop.get("video"))
        loop["audioSync"] = True
        if essay and essay.get("leadingQuestion"):
            loop["leadingQuestion"] = essay["leadingQuestion"]

        # Spot-check keeper casing
        if "keeper" in (loop.get("prayer") or "").lower():
            m = re.search(r".{0,20}keeper.{0,10}", loop["prayer"], re.I)
            if m:
                print(f"    keeper sample: {m.group(0)!r}")

    if closing and payload.get("closing"):
        payload["closing"]["prayer"] = closing

    if saved_wrap is not None:
        payload["wrapFraction"] = saved_wrap

    LOOPS_JSON.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote {LOOPS_JSON}")

    # Sync Excel
    try:
        import openpyxl
    except ImportError:
        print("openpyxl missing; skip excel")
        return 0

    if XLSX.is_file():
        wb = openpyxl.load_workbook(XLSX)
        ws = wb.active
        loop_i = 0
        for row in range(2, ws.max_row + 1):
            if not ws.cell(row, 1).value:
                continue
            leading = ws.cell(row, 2).value
            if not leading:
                if payload.get("closing"):
                    ws.cell(row, 1).value = payload["closing"].get("prayer")
                continue
            if loop_i >= len(loops):
                break
            loop = loops[loop_i]
            ws.cell(row, 1).value = loop.get("prayer")
            ws.cell(row, 2).value = loop.get("leadingQuestion")
            ws.cell(row, 3).value = VIDEO_BY_KEY.get(loop.get("videoKey"), loop.get("videoKey"))
            loop_i += 1
        wb.save(XLSX)
        print(f"Updated {XLSX}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
